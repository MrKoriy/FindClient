"""Table export: styled XLSX (default) and Excel-friendly CSV."""

import csv
import io
from collections.abc import Sequence
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# (attribute, header, column width)
Column = tuple[str, str, int]

ORG_COLUMNS: list[Column] = [
    ("name", "Название", 36),
    ("phone", "Телефон", 22),
    ("email", "Email", 26),
    ("website", "Сайт", 28),
    ("address", "Адрес", 40),
    ("rating", "Рейтинг", 9),
    ("socials", "Соцсети", 30),
    ("reviews", "Отзывов", 9),
    ("branches", "Филиалов", 9),
    ("category", "Рубрика", 26),
    ("city", "Город", 16),
    ("source", "Источник", 10),
    ("score", "Скоринг", 9),
    ("url", "Карточка", 30),
]

TG_LEAD_COLUMNS: list[Column] = [
    ("name", "Имя", 26),
    ("username", "Username", 20),
    ("link", "Ссылка", 28),
    ("phone", "Телефон", 18),
    ("bio", "Био", 40),
    ("messages", "Сообщений", 11),
    ("chats", "Чаты", 30),
    ("sample", "Пример сообщения", 60),
    ("last_seen", "Последнее сообщение", 20),
    ("score", "Скоринг", 9),
]

ORDER_COLUMNS: list[Column] = [
    ("source", "Источник", 12),
    ("title", "Заказ", 50),
    ("budget", "Бюджет", 16),
    ("published", "Опубликован", 20),
    ("url", "Ссылка", 40),
    ("description", "Описание", 80),
]

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_ZEBRA_FILL = PatternFill("solid", fgColor="F2F6FC")


def _cell(item: Any, attr: str) -> Any:
    value = getattr(item, attr, "") if not isinstance(item, dict) else item.get(attr, "")
    if value is None:
        return ""
    if isinstance(value, (set, list, tuple)):
        return ", ".join(sorted(str(v) for v in value))
    return value


def to_csv(items: Sequence[Any], columns: list[Column]) -> bytes:
    """CSV with utf-8-sig BOM and ';' so Russian Excel opens it in columns."""
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow([c[1] for c in columns])
    for item in items:
        writer.writerow([str(_cell(item, c[0])) for c in columns])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(items: Sequence[Any], columns: list[Column], title: str = "Лиды") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Лиды"

    ws.append([c[1] for c in columns])
    for idx, (_, _, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width

    link_cols = {i for i, c in enumerate(columns, 1) if c[0] in ("url", "link", "website")}
    for row_idx, item in enumerate(items, 2):
        for col_idx, (attr, _, _) in enumerate(columns, 1):
            value = _cell(item, attr)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx in link_cols and isinstance(value, str) and value:
                first = value.split(",")[0].strip()
                if not first.startswith(("http://", "https://", "tg://")):
                    first = "https://" + first
                cell.hyperlink = first
                cell.font = Font(color="0563C1", underline="single")
            if row_idx % 2 == 0:
                cell.fill = _ZEBRA_FILL
            cell.alignment = Alignment(vertical="top", wrap_text=attr in ("description", "sample", "bio"))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(items) + 1)}"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export(items: Sequence[Any], columns: list[Column], fmt: str = "xlsx", title: str = "Лиды") -> tuple[bytes, str]:
    """Return (file bytes, extension)."""
    if fmt == "csv":
        return to_csv(items, columns), "csv"
    return to_xlsx(items, columns, title=title), "xlsx"
