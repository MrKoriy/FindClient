"""Tests for table export, niche catalog, Telegram helpers and new DB features."""

import io

import pytest
from openpyxl import load_workbook

from data.niches import CATEGORIES, NICHES, get_niche, niches_in
from db.database import Database
from models.organization import Organization
from models.tg_lead import TgLead
from services.export import ORG_COLUMNS, TG_LEAD_COLUMNS, to_csv, to_xlsx
from services.telegram_service import PHONE_RE, SITE_REQUEST_RE, _business_score, _clean_phone


class TestExport:

    def test_xlsx_has_headers_links_and_filter(self):
        orgs = [Organization(id="1", name="Ромашка", website="romashka.ru", url="https://2gis.ru/moscow/firm/1")]
        wb = load_workbook(io.BytesIO(to_xlsx(orgs, ORG_COLUMNS)))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        assert headers[:4] == ["Название", "Телефон", "Email", "Сайт"]
        assert ws.cell(row=2, column=1).value == "Ромашка"
        assert ws.cell(row=2, column=4).hyperlink.target == "https://romashka.ru"
        assert ws.freeze_panes == "A2" and ws.auto_filter.ref

    def test_csv_semicolon_bom_and_sets(self):
        lead = TgLead(user_id=1, username="ivan", chats={"@b", "@a"})
        data = to_csv([lead], TG_LEAD_COLUMNS)
        text = data.decode("utf-8-sig")
        assert data.startswith(b"\xef\xbb\xbf")
        assert "https://t.me/ivan" in text and "@a, @b" in text and ";" in text
        assert "None" not in text


class TestNiches:

    def test_catalog_consistent(self):
        ids = [n.id for n in NICHES]
        assert len(ids) == len(set(ids)) >= 30
        assert all(n.queries for n in NICHES)
        assert all(len(f"sc:n:{n.id}".encode()) <= 64 for n in NICHES)  # Telegram callback_data limit
        assert sum(len(niches_in(c)) for c in CATEGORIES) == len(NICHES)
        assert get_niche("builders_brigades").map_presence == "no"
        assert get_niche("builders_brigades").tg_chats


class TestTelegramHelpers:

    def test_phone_regex_and_clean(self):
        found = PHONE_RE.findall("Звоните 8 (916) 123-45-67 или +7 999 000 11 22")
        assert [_clean_phone(p) for p in found] == ["+79161234567", "+79990001122"]

    def test_site_request_regex(self):
        assert SITE_REQUEST_RE.search("Добрый день, нужен сайт для бригады")
        assert SITE_REQUEST_RE.search("Кто может сделать лендинг недорого?")
        assert not SITE_REQUEST_RE.search("Нужен бетон М300 на завтра")

    def test_business_score(self):
        assert _business_score("Бригада, выполним ремонт под ключ, договор", 5, True) > _business_score("привет", 1, False)


@pytest.mark.asyncio
async def test_db_settings_orders_and_migration(tmp_path):
    path = str(tmp_path / "m.db")
    import aiosqlite
    # Old schema from the first release (no city/source columns).
    async with aiosqlite.connect(path) as old:
        await old.executescript("""
            CREATE TABLE scrape_sessions (id INTEGER PRIMARY KEY AUTOINCREMENT, niche TEXT NOT NULL,
                count INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL DEFAULT (datetime('now')));
            CREATE TABLE organizations (org_id TEXT NOT NULL, session_id INTEGER NOT NULL,
                name TEXT NOT NULL DEFAULT '', phone TEXT NOT NULL DEFAULT '', email TEXT NOT NULL DEFAULT '',
                website TEXT NOT NULL DEFAULT '', address TEXT NOT NULL DEFAULT '', rating REAL NOT NULL DEFAULT 0.0,
                socials TEXT NOT NULL DEFAULT '', PRIMARY KEY (org_id, session_id));
            INSERT INTO scrape_sessions (niche, count) VALUES ('кафе', 1);
            INSERT INTO organizations (org_id, session_id, name) VALUES ('1', 1, 'Old');
        """)
        await old.commit()

    db = Database(path)
    await db.connect()
    assert await db.get_known_org_ids("кафе") == {"1"}
    sid = await db.save_session("кафе", [{"id": "2", "name": "New", "city": "Казань", "score": 50}], city="Казань")
    rows = await db.get_session_orgs(sid)
    assert rows[0]["city"] == "Казань" and rows[0]["score"] == 50

    await db.set_setting(5, "k", ["a", "b"])
    assert await db.get_setting(5, "k") == ["a", "b"]
    assert await db.get_setting(5, "missing", 1) == 1
    await db.set_setting(5, "on", True)
    assert await db.chats_with_setting("on", True) == [5]
    assert await db.all_values("k") == [["a", "b"]]

    assert await db.filter_new_order_uids(["x:1", "x:2"]) == {"x:1", "x:2"}
    await db.mark_orders_seen([{"uid": "x:1", "source": "x", "matched": True}])
    assert await db.filter_new_order_uids(["x:1", "x:2"]) == {"x:2"}

    await db.save_tg_leads([{"user_id": 10, "username": "u", "chats": {"@c"}}], niche="стройка")
    assert await db.get_known_tg_user_ids() == {10}
    await db.close()
